import re
import pandas as pd
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# ============================================================
# INPUT
# ============================================================

log_text = r"""

"""

# ============================================================
# CONFIGURATION
# ============================================================

# Recording mode:
#   "step"  -> record every STEP_INTERVAL steps
#   "epoch" -> record only the first status line after each epoch marker
RECORD_MODE = "step"

STEP_INTERVAL = 25

EXCEL_FILE = Path(r"F:\StableDiffusion\Datasets\Z Logs\ParsedLogs.xlsx")

# ============================================================
# REGEX PATTERNS
# ============================================================

epoch_pattern = re.compile(
    r"epoch\s+(?P<current>\d+)/(?P<total>\d+)",
    re.IGNORECASE,
)

step_pattern = re.compile(
    r"steps:.*?(?P<step>\d+)/(?P<total>\d+)"
)

loss_pattern = re.compile(
    r"avr_loss=(?P<loss>[0-9]*\.?[0-9]+)"
)

speed_pattern = re.compile(
    r"(?P<speed>[0-9]*\.?[0-9]+)\s*(?P<unit>s/it|it/s)"
)

model_pattern = re.compile(
    r"saving checkpoint:.*?[\\/](?P<model>[^\\/]+)-step\d+\.safetensors",
    re.IGNORECASE,
)

# ============================================================
# PRE-SCAN LOG
# ============================================================

# Model name
model_match = model_pattern.search(log_text)
model_name = model_match.group("model") if model_match else None

# Total epochs (epoch 1 is never printed)
epoch_match = epoch_pattern.search(log_text)
total_epochs = int(epoch_match.group("total")) if epoch_match else None

# ============================================================
# PARSE LOG
# ============================================================

records = {}

current_epoch = 1
record_next_epoch_line = False

for line in log_text.splitlines():

    # --------------------------------------------------------
    # Epoch marker
    # --------------------------------------------------------

    epoch_match = epoch_pattern.search(line)

    if epoch_match:
        current_epoch = int(epoch_match.group("current"))
        record_next_epoch_line = True
        continue

    # --------------------------------------------------------
    # Step line
    # --------------------------------------------------------

    step_match = step_pattern.search(line)

    if not step_match:
        continue

    loss_match = loss_pattern.search(line)
    speed_match = speed_pattern.search(line)

    if not (loss_match and speed_match):
        continue

    step = int(step_match.group("step"))
    total_steps = int(step_match.group("total"))

    loss = float(loss_match.group("loss"))

    speed = float(speed_match.group("speed"))
    speed_unit = speed_match.group("unit")

    # --------------------------------------------------------
    # Decide whether to record
    # --------------------------------------------------------

    keep = False

    if RECORD_MODE.lower() == "step":

        if step % STEP_INTERVAL == 0:
            keep = True

    elif RECORD_MODE.lower() == "epoch":

        if record_next_epoch_line:
            keep = True
            record_next_epoch_line = False

    else:
        raise ValueError("RECORD_MODE must be 'step' or 'epoch'.")

    # --------------------------------------------------------
    # Record the data for the current step/epoch
    # --------------------------------------------------------

    if keep:

        # Use step number or epoch number as the unique key
        key = current_epoch
        if RECORD_MODE.lower() == "step":
            key = step

        # Later occurrences overwrite earlier ones
        records[key] = {
            #"Epoch": current_epoch,
            "Step": step,
            "Loss": loss,
            #"Speed": speed,
            #"Speed Unit": speed_unit,
        }

# ============================================================
# OUTPUT
# ============================================================

df = pd.DataFrame(records.values())

print("\nModel:", model_name if model_name else "(not found)")
print(df)

# Uncomment if desired
# if model_name:
#     df.to_csv(f"{model_name}_loss.csv", index=False)
# else:
#     df.to_csv("training_loss.csv", index=False)

def RecordLogs():
    # --------------------------------------------------------
    # Create or open workbook
    # --------------------------------------------------------

    if EXCEL_FILE.exists():
        wb = load_workbook(EXCEL_FILE)
    else:
        wb = Workbook()

        # Remove default empty sheet
        ws = wb.active
        if ws is not None and ws.title == "Sheet":
            wb.remove(ws)

    # --------------------------------------------------------
    # Determine sheet name
    # --------------------------------------------------------

    base_name = model_name if model_name else "Unknown Model"

    sheet_name = base_name
    counter = 1

    while sheet_name in wb.sheetnames:
        sheet_name = f"{base_name} ({counter})"
        counter += 1

    # --------------------------------------------------------
    # Create worksheet
    # --------------------------------------------------------

    ws = wb.create_sheet(title=sheet_name)

    # Write Model Name
    ws.append([f"Model:", model_name if model_name else "(not found)"])

    # Write Generated Timestamp
    ws.append([f"Generated on {datetime.now():%Y-%m-%d %H:%M:%S}"])

    # Write dataframe
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)

    # Optional: Freeze header
    ws.freeze_panes = "A4"

    # Save workbook
    wb.save(EXCEL_FILE)

    print(f"\nSaved worksheet '{sheet_name}' to:")
    print(EXCEL_FILE)

if not df.empty:
    RecordLogs()
else:
    print("\nNo records to save.")